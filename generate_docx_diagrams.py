from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


BASE = Path(".")
ASSETS = BASE / "docx_assets"
ASSETS.mkdir(exist_ok=True)
DOCX_PATH = BASE / "smart_calendar_coursework.docx"

try:
    FONT = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 24)
    FONT_B = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 24)
    FONT_S = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 18)
except Exception:
    FONT = ImageFont.load_default()
    FONT_B = FONT
    FONT_S = FONT

BG = (255, 255, 255)
BOX = (239, 246, 255)
LINE = (30, 64, 175)
TEXT = (15, 23, 42)
BORDER = (37, 99, 235)


def wrap_text(draw, text, max_w, fnt):
    words = text.split()
    lines = []
    cur = ""
    for word in words:
        test = (cur + " " + word).strip()
        if draw.textbbox((0, 0), test, font=fnt)[2] <= max_w or not cur:
            cur = test
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def box(draw, xy, title, subtitle=None):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=16, fill=BOX, outline=BORDER, width=3)
    lines = [title] + ([subtitle] if subtitle else [])
    y = y1 + ((y2 - y1) - len(lines) * 27) // 2
    for index, line in enumerate(lines):
        font = FONT_B if index == 0 else FONT_S
        for piece in wrap_text(draw, line, (x2 - x1) - 18, font):
            bb = draw.textbbox((0, 0), piece, font=font)
            draw.text((x1 + (x2 - x1 - bb[2]) // 2, y), piece, font=font, fill=TEXT)
            y += 25


def arrow(draw, a, b, label=None):
    import math

    draw.line([a, b], fill=LINE, width=4)
    angle = math.atan2(b[1] - a[1], b[0] - a[0])
    length = 14
    points = []
    for delta in (2.6, -2.6):
        points.append(
            (
                b[0] - length * math.cos(angle + delta),
                b[1] - length * math.sin(angle + delta),
            )
        )
    draw.polygon([b, points[0], points[1]], fill=LINE)
    if label:
        mx = (a[0] + b[0]) // 2
        my = (a[1] + b[1]) // 2 - 26
        bb = draw.textbbox((0, 0), label, font=FONT_S)
        draw.rectangle((mx - bb[2] // 2 - 6, my - 3, mx + bb[2] // 2 + 6, my + 22), fill=BG)
        draw.text((mx - bb[2] // 2, my), label, font=FONT_S, fill=TEXT)


def save_arch(path):
    image = Image.new("RGB", (1000, 650), BG)
    draw = ImageDraw.Draw(image)
    box(draw, (300, 40, 700, 140), "Frontend", "Next.js, React, TypeScript")
    box(draw, (300, 250, 700, 350), "Backend", "FastAPI, Pydantic, SQLAlchemy")
    box(draw, (80, 500, 420, 600), "PostgreSQL", "users, events, notes")
    box(draw, (580, 500, 920, 600), "Ollama", "локальная LLM")
    arrow(draw, (500, 140), (500, 250), "HTTP/JSON")
    arrow(draw, (390, 350), (250, 500), "ORM")
    arrow(draw, (610, 350), (750, 500), "OpenAI-compatible API")
    image.save(path)


def save_model(path):
    image = Image.new("RGB", (1000, 620), BG)
    draw = ImageDraw.Draw(image)
    box(draw, (60, 220, 310, 340), "User", "id, email, password_hash")
    box(draw, (400, 60, 700, 180), "Event", "user_id, title, date, type")
    box(draw, (400, 250, 700, 370), "Note", "user_id, title, text, date")
    box(draw, (400, 440, 760, 560), "EventTemplate", "user_id, title, duration")
    arrow(draw, (310, 280), (400, 120), "1:N")
    arrow(draw, (310, 280), (400, 310), "1:N")
    arrow(draw, (310, 280), (400, 500), "1:N")
    image.save(path)


def save_login(path):
    image = Image.new("RGB", (1000, 430), BG)
    draw = ImageDraw.Draw(image)
    box(draw, (60, 150, 280, 250), "Frontend")
    box(draw, (390, 150, 610, 250), "FastAPI")
    box(draw, (720, 150, 940, 250), "PostgreSQL")
    arrow(draw, (280, 180), (390, 180), "email, password")
    arrow(draw, (610, 180), (720, 180), "поиск user")
    arrow(draw, (720, 220), (610, 220), "password_hash")
    arrow(draw, (390, 220), (280, 220), "access + refresh")
    image.save(path)


def save_api_layer(path):
    image = Image.new("RGB", (1000, 560), BG)
    draw = ImageDraw.Draw(image)
    box(draw, (330, 40, 670, 130), "React-страница")
    box(draw, (330, 210, 670, 310), "src/lib/api.ts")
    box(draw, (330, 400, 670, 500), "FastAPI backend")
    box(draw, (730, 210, 950, 310), "Токены", "access / refresh")
    arrow(draw, (500, 130), (500, 210), "вызов функции")
    arrow(draw, (500, 310), (500, 400), "HTTP")
    arrow(draw, (670, 260), (730, 260), "чтение/запись")
    arrow(draw, (330, 450), (180, 450))
    arrow(draw, (180, 450), (180, 85))
    arrow(draw, (180, 85), (330, 85), "JSON")
    image.save(path)


def save_ai(path):
    image = Image.new("RGB", (1000, 700), BG)
    draw = ImageDraw.Draw(image)
    box(draw, (330, 30, 670, 120), "Страница аналитики")
    box(draw, (330, 170, 670, 260), "FastAPI /api/ai")
    box(draw, (330, 320, 670, 430), "AI service", "формирование промпта")
    box(draw, (330, 500, 670, 600), "Ollama", "локальная LLM")
    box(draw, (730, 320, 960, 430), "PostgreSQL", "events, notes")
    arrow(draw, (500, 120), (500, 170), "запрос")
    arrow(draw, (500, 260), (500, 320))
    arrow(draw, (670, 375), (730, 375), "данные")
    arrow(draw, (500, 430), (500, 500), "prompt")
    arrow(draw, (330, 550), (170, 550))
    arrow(draw, (170, 550), (170, 75))
    arrow(draw, (170, 75), (330, 75), "отчет")
    image.save(path)


def save_refresh(path):
    image = Image.new("RGB", (1000, 700), BG)
    draw = ImageDraw.Draw(image)
    box(draw, (330, 30, 670, 120), "Запрос к API")
    box(draw, (330, 170, 670, 260), "Проверка access-токена")
    box(draw, (70, 350, 390, 450), "Токен действителен", "данные возвращены")
    box(draw, (610, 350, 930, 450), "Токен истек", "ошибка 401")
    box(draw, (610, 500, 930, 590), "Запрос refresh")
    arrow(draw, (500, 120), (500, 170))
    arrow(draw, (400, 260), (260, 350), "да")
    arrow(draw, (600, 260), (760, 350), "нет")
    arrow(draw, (770, 450), (770, 500))
    arrow(draw, (610, 545), (500, 545))
    arrow(draw, (500, 545), (500, 120), "новый access")
    image.save(path)


DIAGRAMS = [
    ("Общая архитектура приложения", "architecture.png", save_arch),
    ("Упрощенная модель данных приложения", "data_model.png", save_model),
    ("Схема входа пользователя", "login_flow.png", save_login),
    ("Слой взаимодействия frontend с backend", "api_layer.png", save_api_layer),
    ("Взаимодействие с локальной нейросетью", "ai_flow.png", save_ai),
    ("Обновление access-токена через refresh-токен", "refresh_flow.png", save_refresh),
]

paths = {}
for label, filename, builder in DIAGRAMS:
    path = ASSETS / filename
    builder(path)
    paths[label] = path

doc = Document(DOCX_PATH)
for paragraph in list(doc.paragraphs):
    text = paragraph.text.strip()
    for label, path in paths.items():
        if label in text:
            if text != f"Рисунок: {label}":
                paragraph.text = f"Рисунок: {label}"
            image_paragraph = paragraph.insert_paragraph_before()
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.add_run().add_picture(str(path), width=Cm(15.5))
            break

doc.save(DOCX_PATH)
print(f"saved {DOCX_PATH.resolve()}")
print(f"images {len(paths)}")
