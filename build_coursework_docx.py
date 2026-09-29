from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(".")
TEX_PATH = ROOT / "smart_calendar_coursework.tex"
DOCX_PATH = ROOT / "smart_calendar_coursework.docx"
ASSETS = ROOT / "docx_assets"


def set_run_font(run, size=14, bold=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def setup_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.different_first_page_header_footer = True

    styles = doc.styles
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)

    normal = styles["Normal"]
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)

    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.bold = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.append(field)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_centered(doc, text="", size=14, bold=False, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)
    return p


def add_left_block(doc, lines, left_indent_cm=9.0, before=0):
    for index, text in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(left_indent_cm)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(before if index == 0 else 0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run)


def add_title_page(doc):
    add_centered(doc, "МИНИСТЕРСТВО ОБРАЗОВАНИЯ РЕСПУБЛИКИ БЕЛАРУСЬ")
    add_centered(doc, "БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ")
    add_centered(doc, "ФАКУЛЬТЕТ ПРИКЛАДНОЙ МАТЕМАТИКИ И ИНФОРМАТИКИ")
    add_centered(doc, "Кафедра многопроцессорных систем и сетей")

    for _ in range(6):
        add_centered(doc)

    add_centered(doc, "РЕАЛИЗАЦИЯ WEB-ПРИЛОЖЕНИЯ", bold=True)
    add_centered(doc, "«УМНЫЙ КАЛЕНДАРЬ»", bold=True)

    for _ in range(2):
        add_centered(doc)

    add_centered(doc, "Курсовой проект")

    for _ in range(4):
        add_centered(doc)

    add_left_block(
        doc,
        [
            "Шибитова Николая Дмитриевича",
            "обучающегося 3 курса специальности",
            "«Прикладная информатика»",
            "",
            "Научный руководитель:",
            "заведующий кафедрой многопроцессорных систем и сетей,",
            "доцент, кандидат физико-математических наук",
            "Андрушкевич А.Н.",
        ],
        left_indent_cm=8.2,
    )

    for _ in range(5):
        add_centered(doc)

    add_centered(doc, "Минск, 2025")
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("СОДЕРЖАНИЕ")
    set_run_font(run, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Оглавление обновится автоматически в Microsoft Word: щелкните правой кнопкой мыши и выберите «Обновить поле»."
    r.append(t)
    fld.append(r)
    p._p.append(fld)

    doc.add_page_break()


CMD_ARG_RE = re.compile(r"\\(textbf|texttt|underline|url|emph)\{([^{}]*)\}")


def clean_latex(text):
    text = text.strip()
    text = text.replace("\\_", "_")
    text = text.replace("\\{", "{").replace("\\}", "}")
    text = text.replace("\\%", "%").replace("\\&", "&")
    text = text.replace("``", "«").replace("''", "»")
    text = text.replace("---", "—").replace("--", "—")
    text = text.replace("~", " ")
    text = re.sub(r"\\addcontentsline\{[^}]+\}\{[^}]+\}\{[^}]+\}", "", text)
    text = re.sub(r"\\(thispagestyle|pagenumbering|setcounter)\{[^}]*\}(\{[^}]*\})?", "", text)
    text = re.sub(r"\\vspace\{[^}]*\}", "", text)
    text = re.sub(r"\\noindent\s*", "", text)
    text = re.sub(r"\\item\s*", "", text)
    while True:
        next_text = CMD_ARG_RE.sub(lambda m: m.group(2), text)
        if next_text == text:
            break
        text = next_text
    text = re.sub(r"\\[a-zA-Z]+\*?(\[[^]]*\])?", "", text)
    text = text.replace("\\\\", "\n")
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def extract_arg(line):
    start = line.find("{")
    if start < 0:
        return ""
    depth = 0
    out = []
    for char in line[start + 1 :]:
        if char == "{":
            depth += 1
            out.append(char)
        elif char == "}":
            if depth == 0:
                break
            depth -= 1
            out.append(char)
        else:
            out.append(char)
    return clean_latex("".join(out))


def add_heading(doc, text, level=1, centered=False):
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 3)}"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, bold=True)


def add_body_paragraph(doc, text):
    for part in clean_latex(text).split("\n"):
        part = part.strip()
        if part:
            p = doc.add_paragraph(part)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def set_cell_text(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=12)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows, caption=None):
    if caption:
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_i, row in enumerate(rows):
        for c_i in range(cols):
            set_cell_text(table.rows[r_i].cells[c_i], row[c_i] if c_i < len(row) else "")


def parse_body(doc):
    text = TEX_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    in_doc = False
    body_started = False
    pending = []
    i = 0

    def flush():
        nonlocal pending
        if pending:
            add_body_paragraph(doc, " ".join(pending))
            pending = []

    while i < len(lines):
        line = lines[i].strip()
        i += 1

        if not in_doc:
            if line == r"\begin{document}":
                in_doc = True
            continue
        if line == r"\end{document}":
            break

        if not body_started:
            if line.startswith(r"\chapter*{ПЕРЕЧЕНЬ"):
                body_started = True
            else:
                continue

        if not line or line.startswith("%") or line == r"\tableofcontents":
            flush()
            continue

        if line.startswith(r"\begin{Verbatim}"):
            flush()
            block = []
            while i < len(lines) and not lines[i].strip().startswith(r"\end{Verbatim}"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            for index, raw in enumerate(block):
                if index:
                    p.add_run("\n")
                run = p.add_run(raw)
                set_run_font(run, size=10, name="Courier New")
            continue

        if line.startswith(r"\begin{longtable}"):
            flush()
            table_lines = []
            caption = None
            while i < len(lines) and not lines[i].strip().startswith(r"\end{longtable}"):
                table_line = lines[i].strip()
                match = re.match(r"\\caption\{(.+)\}\\\\", table_line)
                if match:
                    caption = clean_latex(match.group(1))
                else:
                    table_lines.append(table_line)
                i += 1
            i += 1
            rows = []
            acc = ""
            for table_line in table_lines:
                if (
                    not table_line
                    or table_line == r"\hline"
                    or table_line.startswith(r"\toprule")
                    or table_line.startswith(r"\midrule")
                    or table_line.startswith(r"\bottomrule")
                ):
                    continue
                acc += " " + table_line
                if table_line.endswith(r"\\"):
                    row = acc.strip()[:-2].strip()
                    acc = ""
                    if row and not row.startswith("\\"):
                        cells = [clean_latex(cell) for cell in row.split("&")]
                        if any(cells):
                            rows.append(cells)
            add_table(doc, rows, caption)
            continue

        if line.startswith(r"\begin{tikzpicture}"):
            while i < len(lines) and not lines[i].strip().startswith(r"\end{tikzpicture}"):
                i += 1
            i += 1
            continue

        if line.startswith(r"\caption"):
            flush()
            caption = extract_arg(line)
            if caption:
                insert_diagram_if_exists(doc, caption)
                p = doc.add_paragraph(f"Рисунок: {caption}")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
            continue

        if line.startswith(r"\chapter"):
            flush()
            title = extract_arg(line)
            if title:
                add_heading(doc, title, 1, True)
            continue

        if line.startswith(r"\section"):
            flush()
            title = extract_arg(line)
            if title:
                add_heading(doc, title, 2)
            continue

        if line.startswith(r"\subsection"):
            flush()
            title = extract_arg(line)
            if title:
                add_heading(doc, title, 3)
            continue

        if line.startswith(r"\begin{") or line.startswith(r"\end{"):
            continue

        pending.append(line)

    flush()


DIAGRAM_FILES = {
    "Общая архитектура приложения": "architecture.png",
    "Упрощенная модель данных приложения": "data_model.png",
    "Схема входа пользователя": "login_flow.png",
    "Слой взаимодействия frontend с backend": "api_layer.png",
    "Взаимодействие с локальной нейросетью": "ai_flow.png",
    "Обновление access-токена через refresh-токен": "refresh_flow.png",
}


def insert_diagram_if_exists(doc, caption):
    filename = DIAGRAM_FILES.get(caption)
    if not filename:
        return
    path = ASSETS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Cm(15.5))


def main():
    doc = Document()
    setup_document(doc)
    add_title_page(doc)
    add_toc(doc)
    parse_body(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH.resolve())


if __name__ == "__main__":
    main()
