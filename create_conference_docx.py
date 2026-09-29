from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt


BASE_DIR = Path(__file__).parent
MD_PATH = BASE_DIR / "conference_doklad.md"
OUT_PATH = BASE_DIR / "conference_doklad.docx"


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)


def add_plain_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text.replace("**", ""))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Mm(7)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)


def build_docx() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line:
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("**") and line.endswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line.strip("*"))
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(line[2:])
        elif len(line) > 3 and line[0].isdigit() and ". " in line[:5]:
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(line.split(". ", 1)[1])
        else:
            add_plain_paragraph(doc, line)

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_docx()
