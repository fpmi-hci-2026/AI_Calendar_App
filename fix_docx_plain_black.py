from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DOCX_PATH = "smart_calendar_coursework.docx"
OUT_PATH = "smart_calendar_coursework_black_plain.docx"


def set_run_black(run):
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if run.font.size is None:
        run.font.size = Pt(14)


def clear_outline_level(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is not None:
        ppr.remove(outline)


def add_tc_field_before(paragraph, title, level):
    tc = OxmlElement("w:fldSimple")
    safe_title = title.replace('"', "'")
    tc.set(qn("w:instr"), f'TC "{safe_title}" \\l {level}')
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    vanish = OxmlElement("w:vanish")
    rpr.append(vanish)
    run.append(rpr)
    tc.append(run)
    paragraph._p.insert(0, tc)


def replace_toc_field(paragraph):
    paragraph.text = ""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r"TOC \f \h \z")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Оглавление обновится автоматически: выделите документ Ctrl+A и нажмите F9."
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def normalize_heading(paragraph):
    text = paragraph.text.strip()
    if not text:
        return
    old_style = paragraph.style.name
    if old_style == "Heading 1":
        level = 1
    elif old_style == "Heading 2":
        level = 2
    elif old_style == "Heading 3":
        level = 3
    else:
        return

    add_tc_field_before(paragraph, text, level)
    paragraph.style = "Normal"
    clear_outline_level(paragraph)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        set_run_black(run)
        run.bold = True


def main():
    doc = Document(DOCX_PATH)

    # Replace TOC field from heading-based TOC to TC-field-based TOC.
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "СОДЕРЖАНИЕ":
            for next_paragraph in doc.paragraphs[index + 1 : index + 4]:
                replace_toc_field(next_paragraph)
                break
            break

    for paragraph in doc.paragraphs:
        normalize_heading(paragraph)
        clear_outline_level(paragraph)
        for run in paragraph.runs:
            set_run_black(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    clear_outline_level(paragraph)
                    for run in paragraph.runs:
                        set_run_black(run)

    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            for paragraph in part.paragraphs:
                clear_outline_level(paragraph)
                for run in paragraph.runs:
                    set_run_black(run)

    try:
        doc.save(DOCX_PATH)
        print("saved", DOCX_PATH)
    except PermissionError:
        doc.save(OUT_PATH)
        print("saved", OUT_PATH)


if __name__ == "__main__":
    main()
